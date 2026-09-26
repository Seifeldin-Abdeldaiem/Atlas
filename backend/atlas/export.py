"""The cleaned file (milestone 4): the customer's own rows and columns,
untouched, with Atlas's findings added as columns on the right, in the format
they uploaded (CSV, Excel or JSON).

Safety:
- Spreadsheet formula injection: text starting with = + - @ tab or CR is
  written as text (Excel) or prefixed with ' (CSV), so opening the file can
  never run a formula. Plain numbers such as -5 are left alone.
- The file is written in a child process with a memory cap and a time limit,
  stored privately under a random key, downloaded only through the API and
  deleted after the retention period or with its dataset.
"""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import json
import logging
import re
from typing import Any
from uuid import UUID, uuid4

import psycopg
from psycopg.types.json import Jsonb

from . import datasets, db, jobs, storage
from .config import get_settings
from .errors import AtlasError

log = logging.getLogger("atlas.export")

ADDED = [
    "atlas_group",
    "atlas_role",
    "atlas_master_code",
    "atlas_confidence",
    "atlas_reason",
    "atlas_lookalike_of",
    "atlas_name_standard",
    "atlas_brand",
    "atlas_part_number",
    "atlas_variant",
    "atlas_group_stock",
    "atlas_group_value",
]
CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "json": "application/json",
}
_RISKY_START = ("=", "+", "-", "@", "\t", "\r")
_PLAIN_NUMBER = re.compile(r"-?(0|[1-9]\d*)(\.\d+)?")


def is_risky(value: str) -> bool:
    return value.startswith(_RISKY_START) and not _PLAIN_NUMBER.fullmatch(value)


def csv_safe(value: object) -> str:
    text = "" if value is None else str(value)
    return "'" + text if is_risky(text) else text


# ---------------------------------------------------------------- rows

def added_columns(
    rows: list[dict],
    mapping: dict,
    groups: list[dict],
    lookalikes: list[tuple[int, int]],
    review_rows: set[int],
) -> dict[int, dict[str, Any]]:
    """What Atlas adds to each row, keyed by row number. rows are
    {"row_number", "original", "norm", "malformed"}."""
    code_col = mapping.get("item_code")
    by_row = {r["row_number"]: r for r in rows}

    def code_of(row_number: int) -> str:
        r = by_row.get(row_number)
        value = (r["original"] or {}).get(code_col) if r and code_col else None
        return value or f"row {row_number}"

    member_of: dict[int, dict] = {}
    for g in groups:
        for m in g["members"]:
            member_of[m["row_number"]] = {**g, "role": m["role"]}
    partners: dict[int, list[int]] = {}
    for a, b in lookalikes:
        partners.setdefault(a, []).append(b)
        partners.setdefault(b, []).append(a)

    out: dict[int, dict[str, Any]] = {}
    for r in rows:
        n, norm = r["row_number"], r["norm"] or {}
        if r["malformed"] or not r["norm"]:
            out[n] = {"atlas_role": "not_analysed"}
            continue
        g = member_of.get(n)
        part = (norm.get("part_base") or "").upper()
        added: dict[str, Any] = {
            "atlas_group": g["label"] if g else "",
            "atlas_role": g["role"] if g else ("needs_review" if n in review_rows else "unique"),
            "atlas_master_code": code_of(g["master_row"]) if g else "",
            "atlas_confidence": g["confidence"] if g else "",
            "atlas_reason": (g["reasons"] or [""])[0] if g else "",
            "atlas_lookalike_of": ", ".join(code_of(p) for p in sorted(partners.get(n, []))[:3]),
            "atlas_name_standard": g["name_standard"] if g else "",
            "atlas_brand": (norm.get("brand") or "").upper(),
            "atlas_part_number": part,
            "atlas_variant": "/".join(v.upper() for v in norm.get("variants", [])),
            "atlas_group_stock": "",
            "atlas_group_value": "",
        }
        if g and g["role"] == "master":
            added["atlas_group_stock"] = g["stock_total"] if g["stock_total"] is not None else ""
            added["atlas_group_value"] = g["value_on_duplicates"] if g["value_on_duplicates"] is not None else ""
        out[n] = added
    return out


def _headers(columns: list[str]) -> list[str]:
    """Added column names, renamed if the file already uses one."""
    taken = set(columns)
    return [name if name not in taken else f"{name} (atlas)" for name in ADDED]


# ---------------------------------------------------------------- writers (run in the sandbox)

def write_file(fmt: str, columns: list[str], rows: list[dict], added: dict[int, dict], summary: dict, groups: list[dict]) -> bytes:
    headers = _headers(columns)
    if fmt == "xlsx":
        return _write_xlsx(columns, headers, rows, added, summary, groups)
    if fmt == "json":
        records = []
        for r in rows:
            extra = added.get(r["row_number"], {})
            records.append({**(r["original"] or {}), **{h: extra.get(k, "") for h, k in zip(headers, ADDED, strict=True)}})
        return json.dumps(records, ensure_ascii=False, indent=1).encode()
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([csv_safe(c) for c in columns + headers])
    for r in rows:
        extra = added.get(r["row_number"], {})
        original = r["original"] or {}
        writer.writerow([csv_safe(original.get(c, "")) for c in columns] + [csv_safe(extra.get(k, "")) for k in ADDED])
    # A byte-order mark makes Excel read UTF-8 (Arabic, £) correctly.
    return buffer.getvalue().encode("utf-8-sig")


def _write_xlsx(columns: list[str], headers: list[str], rows: list[dict], added: dict[int, dict], summary: dict, groups: list[dict]) -> bytes:
    import openpyxl
    from openpyxl.cell import WriteOnlyCell

    book = openpyxl.Workbook(write_only=True)
    sheet = book.create_sheet("Items")

    def cell(value: Any):
        if isinstance(value, (int, float)):
            return value
        text = "" if value is None else str(value)
        if _PLAIN_NUMBER.fullmatch(text) and len(text) < 16:
            return float(text) if "." in text else int(text)
        c = WriteOnlyCell(sheet, value=text)
        c.data_type = "s"  # always text: never evaluated as a formula
        return c

    sheet.append([cell(h) for h in columns + headers])
    for r in rows:
        extra = added.get(r["row_number"], {})
        original = r["original"] or {}
        sheet.append([cell(original.get(c, "")) for c in columns] + [cell(extra.get(k, "")) for k in ADDED])

    info = book.create_sheet("Atlas summary")
    stock = summary.get("stock") or {}
    currency = stock.get("currency") or ""
    for label, value in (
        ("Lines checked", summary.get("items")),
        ("Duplicate groups", summary.get("groups")),
        ("Duplicate lines", summary.get("duplicate_lines")),
        ("Look-alikes kept apart", summary.get("lookalikes")),
        ("Pairs that need your review", summary.get("needs_review")),
        ("Units on duplicate lines", stock.get("units_on_duplicates")),
        (f"Stock value on duplicate lines ({currency})" if currency else "Stock value on duplicate lines", stock.get("value_on_duplicates")),
        ("Groups whose unit costs differ by more than 20%", stock.get("groups_with_cost_gap")),
    ):
        info.append([label, "" if value is None else value])
    info.append([])
    info.append(["Group", "Standard name", "Lines", "Confidence", "Total in stock", "Stock on duplicate lines", "Value on duplicate lines", "Reason"])
    for g in groups:
        info.append([
            cell(g["label"]), cell(g["name_standard"] or ""), len(g["members"]), g["confidence"],
            "" if g["stock_total"] is None else g["stock_total"],
            "" if g["stock_on_duplicates"] is None else g["stock_on_duplicates"],
            "" if g["value_on_duplicates"] is None else g["value_on_duplicates"],
            cell((g["reasons"] or [""])[0]),
        ])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


# ---------------------------------------------------------------- database side

def _num(v: object) -> float | None:
    return None if v is None else float(v)


def collect(conn: psycopg.Connection, dataset: dict) -> tuple[list[dict], list[dict], list[tuple[int, int]], set[int]]:
    rows = conn.execute(
        "SELECT row_number, original, norm, malformed_reason AS malformed FROM tasks WHERE dataset_id = %s ORDER BY row_number",
        (dataset["id"],),
    ).fetchall()
    groups = conn.execute("SELECT * FROM match_groups WHERE dataset_id = %s ORDER BY label", (dataset["id"],)).fetchall()
    members = conn.execute("SELECT group_id, row_number, role FROM match_members WHERE dataset_id = %s", (dataset["id"],)).fetchall()
    by_group: dict = {}
    for m in members:
        by_group.setdefault(m["group_id"], []).append({"row_number": m["row_number"], "role": m["role"]})
    group_list = [
        {
            "label": g["label"], "confidence": g["confidence"], "master_row": g["master_row"], "reasons": g["reasons"],
            "name_standard": g["name_standard"], "stock_total": _num(g["stock_total"]),
            "stock_on_duplicates": _num(g["stock_on_duplicates"]), "value_on_duplicates": _num(g["value_on_duplicates"]),
            "members": sorted(by_group.get(g["id"], []), key=lambda m: m["row_number"]),
        }
        for g in groups
    ]
    pairs = conn.execute("SELECT kind, row_a, row_b FROM match_pairs WHERE dataset_id = %s", (dataset["id"],)).fetchall()
    lookalikes = [(p["row_a"], p["row_b"]) for p in pairs if p["kind"] == "lookalike"]
    review = {r for p in pairs if p["kind"] == "review" for r in (p["row_a"], p["row_b"])}
    return [dict(r) for r in rows], group_list, lookalikes, review


def start(conn: psycopg.Connection, org_id: UUID, dataset: dict) -> dict:
    if dataset["status"] != "analysed":
        raise AtlasError("not_ready", status_code=409)
    discard(conn, dataset["id"])
    state = {"state": "building", "run_id": str(uuid4())}
    conn.execute("UPDATE datasets SET export = %s WHERE id = %s", (Jsonb(state), dataset["id"]))
    jobs.enqueue(conn, org_id, jobs.BUILD_EXPORT, {"dataset_id": str(dataset["id"]), "run_id": state["run_id"]})
    return state


def discard(conn: psycopg.Connection, dataset_id: UUID) -> None:
    """Delete the export file and forget it."""
    for f in conn.execute(
        "SELECT object_key FROM dataset_files WHERE dataset_id = %s AND kind = 'export' AND deleted_at IS NULL", (dataset_id,)
    ).fetchall():
        try:
            storage.delete(f["object_key"])
        except Exception:
            log.exception("could not delete old export", extra={"dataset_id": str(dataset_id)})
    conn.execute("DELETE FROM dataset_files WHERE dataset_id = %s AND kind = 'export'", (dataset_id,))
    conn.execute("UPDATE datasets SET export = NULL WHERE id = %s", (dataset_id,))


def filename(dataset: dict) -> str:
    stem = re.sub(r"[^\w\- ]+", "", dataset["name"] or "export").strip() or "export"
    return f"{stem[:80]} - Atlas.{dataset['file_format'] or 'csv'}"


def run(org_id: UUID, dataset_id: UUID, run_id: str) -> bool:
    """Build and store the file (worker). False if it's no longer wanted."""
    from .sandbox import run_limited

    settings = get_settings()
    with db.tenant(org_id) as conn:
        dataset = datasets.get_dataset(conn, dataset_id)
        if dataset["status"] != "analysed" or (dataset["export"] or {}).get("run_id") != run_id:
            return False
        rows, groups, lookalikes, review = collect(conn, dataset)
    fmt = dataset["file_format"] or "csv"
    added = added_columns(rows, dataset["mapping"], groups, lookalikes, review)
    summary = (dataset["analysis"] or {}).get("summary") or {}
    content = run_limited(write_file, (fmt, list(dataset["columns"]), rows, added, summary, groups), settings.parse_timeout_seconds, settings.parse_memory_mb)

    key = storage.new_key(org_id, dataset_id) + "-export"
    storage.put(key, content)
    expires = dt.datetime.now(dt.UTC) + dt.timedelta(days=settings.raw_file_retention_days)
    with db.tenant(org_id) as conn:
        current = conn.execute("SELECT status, export FROM datasets WHERE id = %s FOR UPDATE", (dataset_id,)).fetchone()
        if current is None or current["status"] != "analysed" or (current["export"] or {}).get("run_id") != run_id:
            storage.delete(key)
            return False
        name = filename(dataset)
        conn.execute(
            """INSERT INTO dataset_files (org_id, dataset_id, kind, object_key, original_filename, size_bytes, sha256, expires_at)
               VALUES (%s, %s, 'export', %s, %s, %s, %s, %s)""",
            (org_id, dataset_id, key, name, len(content), hashlib.sha256(content).hexdigest(), expires),
        )
        state = {"state": "ready", "run_id": run_id, "file": {"name": name, "size_bytes": len(content), "format": fmt, "expires_at": expires.isoformat()}}
        conn.execute("UPDATE datasets SET export = %s WHERE id = %s", (Jsonb(state), dataset_id))
        datasets.audit(conn, org_id, None, "dataset.exported", dataset_id, size_bytes=len(content), format=fmt)
    return True


def fail(org_id: UUID, dataset_id: UUID, run_id: str) -> None:
    with db.tenant(org_id) as conn:
        conn.execute(
            "UPDATE datasets SET export = %s WHERE id = %s AND export->>'run_id' = %s",
            (Jsonb({"state": "failed", "run_id": run_id, "error_code": "export_failed"}), dataset_id, run_id),
        )


def download(conn: psycopg.Connection, dataset: dict) -> tuple[bytes, str, str]:
    """(content, filename, content type) of the ready export."""
    f = conn.execute(
        "SELECT object_key, original_filename FROM dataset_files WHERE dataset_id = %s AND kind = 'export' AND deleted_at IS NULL",
        (dataset["id"],),
    ).fetchone()
    if f is None or (dataset["export"] or {}).get("state") != "ready":
        raise AtlasError("not_found", status_code=404)
    fmt = dataset["file_format"] or "csv"
    return storage.get_any(f["object_key"]), f["original_filename"], CONTENT_TYPES[fmt]
