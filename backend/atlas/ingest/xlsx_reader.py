"""Read Excel workbooks (.xlsx). Opened read-only with cached values only:
formulas are never evaluated, macros are refused earlier by detect.py, and
XML is parsed through defusedxml (openpyxl uses it when installed)."""

from __future__ import annotations

import datetime as dt
import io
import warnings
import zipfile

from ..errors import IngestError
from .detect import check_zip_safety
from .tabular import build_table
from .types import DEFAULT_LIMITS, Limits, SheetInfo, Table


def cell_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return repr(value)
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    return str(value)


def _open_workbook(data: bytes, limits: Limits):
    import openpyxl  # imported lazily so the rest of ingest has no hard dependency

    try:
        check_zip_safety(zipfile.ZipFile(io.BytesIO(data)).infolist(), limits)
    except zipfile.BadZipFile:
        raise IngestError("unreadable_file") from None
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except IngestError:
        raise
    except Exception:  # openpyxl raises many types for damaged files
        raise IngestError("unreadable_file") from None


def _is_visible(ws) -> bool:
    return getattr(ws, "sheet_state", "visible") == "visible"


def list_sheets(workbook) -> list[SheetInfo]:
    sheets = []
    for ws in workbook.worksheets:
        if not _is_visible(ws):
            continue
        try:
            approx = max((ws.max_row or 1) - 1, 0)
        except Exception:
            approx = 0
        sheets.append(SheetInfo(name=ws.title, approx_rows=approx))
    return sheets


def read_xlsx(data: bytes, sheet: str | None = None, limits: Limits = DEFAULT_LIMITS) -> tuple[Table, list[SheetInfo]]:
    workbook = _open_workbook(data, limits)
    try:
        sheets = list_sheets(workbook)
        if not sheets:
            raise IngestError("no_rows")

        if sheet is not None:
            if sheet not in {s.name for s in sheets}:
                raise IngestError("sheet_not_found")
            candidates = [sheet]
        else:
            # The first visible sheet that holds any tasks.
            candidates = [s.name for s in sheets]

        last_error: IngestError | None = None
        for name in candidates:
            ws = workbook[name]
            try:
                table = build_table(_rows(ws), limits, meta={"sheet": name})
                return table, sheets
            except IngestError as error:
                if error.code != "no_rows":
                    raise
                last_error = error
        raise last_error or IngestError("no_rows")
    finally:
        workbook.close()


def _rows(ws):
    try:
        for row_number, row in enumerate(ws.iter_rows(values_only=True), start=1):
            yield row_number, [cell_to_text(v) for v in row]
    except IngestError:
        raise
    except Exception:
        raise IngestError("unreadable_file") from None
